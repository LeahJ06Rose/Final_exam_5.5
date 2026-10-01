import os
import csv
import pandas as pd
import matplotlib.pyplot as plt
import openpyxl
from datetime import datetime

from os import chdir

from openpyxl import load_workbook

from openpyxl.chart import Reference,PieChart

path = "C:\\FinalExam\\final.csv"

def askUser():
    total = 0
    #The loop allows the user to run the code 5 times
    #asking for a number each time.
    for i in range(5):
        num = float(input(f"Please enter a number {i + 1}: "))
        total += num
    print(f"The total of your five numbers is: {total}")

def askIncome():
    #This loop asks the user to input 5 names and
    #their incomes.
    with open("final.csv", "a", newline = "") as f:
        writer = csv.writer(f)

        for i in range(5):
            name = input(f"Enter name #{i+1}: ")
            income = input(f"Enter annual income for {name}: ")
            writer.writerow([name, income])

os.chdir("C:\\FinalExam")

if os.path.exists("final.xlsx"):
    wb=openpyxl.load_workbook("final.xlsx")

else:
    wb = openpyxl.Workbook()
    wb.save("final.xlsx")

ws=wb.active

rows = ws.max_row
cols = ws.max_column


def excelPie():
    df = pd.read_csv("final.csv")
    
    #get the labels
    myDataLabels = Reference(ws,min_col = 1, min_row=2, max_row=rows)

    #get the data
    myData=Reference(ws,min_col=2, min_row=2,max_row=rows)

    #get the value from 'Sheet1' !A1
    xAxisName = ws['A1'].value

    #get the value from 'Sheet1' !B1
    yAxisName = ws['B1'].value

    #create a PieChart object
    myPieChart=PieChart()
    timestamp = datetime.now().strftime("%m-%d-%Y")
    #sets title attribute
    myPieChart.title = f"leajon3478 {timestamp}"
    
    #adds the data to be plotted
    myPieChart.add_data(myData)

    #Adds chart, the cell will appear in upper
    #left corner of plot area
    ws.add_chart(myPieChart,"A10")
    
    #sets the value that will appear
    myPieChart.set_categories(myDataLabels)

    myPieChart.height = 15
    myPieChart.width = 15

    wb.save("final.xlsx")

alphas = []
for charCode in range(65,91):
    alphas.append(chr(charCode))

for row in range(1,rows+1):
    for column in range(cols):
        print(ws[alphas[column]+str(row)].value,end="\t")
    print("\n")

def verticalBar():
    df = pd.read_csv("C:\\FinalExam\\final.csv")
    x = []
    y = []
        
    plt.bar(x,y, color = "g",width = 0.50, label = "Income")
    timestamp = datetime.now().strftime("%m-%d-%Y")
    
    plt.title(f"leajon3478 {timestamp}")

    plt.xticks(rotation = 40)
    plt.xlabel("Name")
    plt.ylabel("Annual Income")
    plt.show()


def main():
    askUser()
    askIncome()
    excelPie()
    verticalBar()
main()
